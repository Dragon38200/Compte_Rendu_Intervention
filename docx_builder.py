# -*- coding: utf-8 -*-
"""
Génération du rapport .docx, portée telle quelle depuis l'application de
bureau (essais_pyqt6_secours.py) mais sans aucune dépendance à Qt : cette
version prend en entrée un simple dictionnaire JSON (le "payload" envoyé
par le frontend web) plus des fichiers (photos / logo / signatures, déjà
reçus en bytes par l'API), et renvoie les bytes du .docx généré.

Toute la logique métier (tableau méta, prestations, galerie photo 2 par 2,
pied de page dynamique avec NUMPAGES, bloc signatures) est identique à la
version bureau, déjà testée et validée avec l'utilisateur.
"""
import os
import io
import copy
import base64
import random
import unicodedata
import re
import tempfile

from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageOps

TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "templates", "Fond de Page Rapport.docx"
)


def slugify(text):
    """Transforme un nom en identifiant de fichier sûr (sans accents/espaces)."""
    text = unicodedata.normalize("NFKD", text or "").encode("ascii", "ignore").decode("ascii")
    text = re.sub(r"[^a-zA-Z0-9]+", "_", text).strip("_").lower()
    return text or "fichier"


def apply_dynamic_footer(doc, title_text):
    """
    Personnalise le pied de page issu du fond de page ('Fond de Page
    Rapport.docx') pour ce rapport précis :
      1. Le champ FILENAME (titre figé du modèle d'origine) est remplacé
         par le titre du rapport en cours (ex: "Client - Équipement").
      2. Le total de pages, codé en dur ("/ 1") dans le modèle, est
         remplacé par un vrai champ NUMPAGES, pour qu'il reste exact même
         si le rapport fait plusieurs pages.
    Opération silencieuse : si la structure attendue n'est pas retrouvée,
    la fonction ne fait rien et le document reste tel quel.
    """
    try:
        for section in doc.sections:
            footer_elm = section.footer._element

            for fld in footer_elm.findall(".//" + qn("w:fldSimple")):
                instr = fld.get(qn("w:instr")) or ""
                if "FILENAME" not in instr:
                    continue
                inner_r = fld.find(qn("w:r"))
                rpr_copy = None
                if inner_r is not None:
                    rpr_el = inner_r.find(qn("w:rPr"))
                    if rpr_el is not None:
                        rpr_copy = copy.deepcopy(rpr_el)
                new_r = OxmlElement("w:r")
                if rpr_copy is not None:
                    new_r.append(rpr_copy)
                new_t = OxmlElement("w:t")
                new_t.set(qn("xml:space"), "preserve")
                new_t.text = title_text
                new_r.append(new_t)
                fld.getparent().replace(fld, new_r)

            for p in footer_elm.findall(".//" + qn("w:p")):
                runs = p.findall(qn("w:r"))
                sep_idx = next(
                    (i for i, r in enumerate(runs)
                     if (r.find(qn("w:t")) is not None and r.find(qn("w:t")).text == " / ")),
                    None
                )
                if sep_idx is None or sep_idx + 1 >= len(runs):
                    continue
                total_run = runs[sep_idx + 1]
                t_el = total_run.find(qn("w:t"))
                if t_el is None or not (t_el.text or "").strip().isdigit():
                    continue

                rpr_el = total_run.find(qn("w:rPr"))

                def make_run(build_fn):
                    r = OxmlElement("w:r")
                    if rpr_el is not None:
                        r.append(copy.deepcopy(rpr_el))
                    build_fn(r)
                    return r

                def add_fldchar(r, type_):
                    fc = OxmlElement("w:fldChar")
                    fc.set(qn("w:fldCharType"), type_)
                    r.append(fc)

                def add_instr(r, text):
                    it = OxmlElement("w:instrText")
                    it.set(qn("xml:space"), "preserve")
                    it.text = text
                    r.append(it)

                def add_text(r, text):
                    t = OxmlElement("w:t")
                    t.set(qn("xml:space"), "preserve")
                    t.text = text
                    r.append(t)

                cached_value = t_el.text
                new_runs = [
                    make_run(lambda r: add_fldchar(r, "begin")),
                    make_run(lambda r: add_instr(r, " NUMPAGES ")),
                    make_run(lambda r: add_fldchar(r, "separate")),
                    make_run(lambda r: add_text(r, cached_value)),
                    make_run(lambda r: add_fldchar(r, "end")),
                ]
                parent = total_run.getparent()
                insert_pos = list(parent).index(total_run)
                for offset, new_r in enumerate(new_runs):
                    parent.insert(insert_pos + offset, new_r)
                parent.remove(total_run)
    except Exception as e:
        print(f"Avertissement : personnalisation du pied de page ignorée ({e})")


def _prepare_image_bytes_for_jpeg(raw_bytes):
    """Ouvre une image (bytes) et renvoie des bytes JPEG utilisables par
    python-docx, en composant les PNG avec transparence sur fond blanc."""
    with Image.open(io.BytesIO(raw_bytes)) as img:
        img = ImageOps.exif_transpose(img)
        width, height = img.size
        if img.mode in ("RGBA", "LA", "P"):
            rgba = img.convert("RGBA")
            background = Image.new("RGB", rgba.size, (255, 255, 255))
            background.paste(rgba, mask=rgba.split()[-1])
            img = background
        elif img.mode != "RGB":
            img = img.convert("RGB")
        buf = io.BytesIO()
        img.save(buf, format="JPEG", quality=90)
        return buf.getvalue(), width, height


def _decode_data_url_or_b64(value):
    """Accepte soit une data-URL ('data:image/png;base64,...'), soit du
    base64 brut, et renvoie les bytes décodés. Renvoie None si vide."""
    if not value:
        return None
    if "," in value and value.strip().lower().startswith("data:"):
        value = value.split(",", 1)[1]
    return base64.b64decode(value)


def build_report_docx(payload, photos, logo_bytes=None):
    """
    Construit le rapport .docx et renvoie ses bytes.

    payload : dict JSON envoyé par le frontend, avec les clés :
        folder, date, technicien, client, site, adresse, contact,
        equipement, serie, observations,
        items: [{designation, qte, pu}, ...],
        content: [
            {"type": "photo", "photo_index": 0, "caption": "...", "scale": 1.0},
            {"type": "text", "text": "..."},
            ...
        ],
        signatures: {
            "technicien": {"nom": "...", "image": "data:image/png;base64,..." | None},
            "client": {...}, "exterieur": {...}
        }

    photos : liste de bytes (une entrée par photo réellement uploadée),
             référencée par "photo_index" dans payload["content"].
    logo_bytes : bytes du logo client (optionnel).
    """
    temp_files = []
    try:
        if os.path.exists(TEMPLATE_PATH):
            doc = Document(TEMPLATE_PATH)
        else:
            doc = Document()

        footer_title_parts = [
            (payload.get("client") or "").strip(),
            (payload.get("equipement") or "").strip(),
        ]
        footer_title = " - ".join(part for part in footer_title_parts if part)
        if footer_title:
            apply_dynamic_footer(doc, footer_title)

        # --- Logo client (optionnel) ---
        if logo_bytes:
            logo_path = os.path.join(tempfile.gettempdir(), f"logo_{os.getpid()}_{random.randint(0, 999999)}.png")
            with open(logo_path, "wb") as f:
                f.write(logo_bytes)
            temp_files.append(logo_path)

            t = doc.add_table(rows=1, cols=1)
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            p = t.rows[0].cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.add_run().add_picture(logo_path, width=Inches(1.8))
            doc.add_paragraph()

        t_title = doc.add_heading("RAPPORT D'INTERVENTION", level=0)
        t_title.alignment = WD_ALIGN_PARAGRAPH.CENTER

        tbl = doc.add_table(rows=9, cols=2)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.style = 'Table Grid'

        meta = [
            ("N° Dossier :", payload.get("folder", "")),
            ("Date :", payload.get("date", "")),
            ("Technicien :", payload.get("technicien", "")),
            ("Client :", payload.get("client", "")),
            ("Site :", payload.get("site", "")),
            ("Adresse Client :", payload.get("adresse", "")),
            ("Contact Client :", payload.get("contact", "")),
            ("Matériel :", payload.get("equipement", "")),
            ("N° de Série :", payload.get("serie", "")),
        ]
        for i, (k, v) in enumerate(meta):
            tbl.rows[i].cells[0].text = k
            tbl.rows[i].cells[1].text = v or "N/C"
            tbl.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True

        doc.add_paragraph()
        doc.add_heading("Observations & Travaux", level=1)
        doc.add_paragraph((payload.get("observations") or "").strip() or "Aucune observation particulière.")

        items = payload.get("items") or []
        if items:
            doc.add_heading("Matériel & Prestations", level=1)
            tbl_item = doc.add_table(rows=len(items) + 2, cols=4)
            tbl_item.alignment = WD_TABLE_ALIGNMENT.CENTER
            tbl_item.style = 'Table Grid'

            headers = ["Désignation", "Qté", "P.U. HT", "Total HT"]
            for i, h in enumerate(headers):
                cell = tbl_item.rows[0].cells[i]
                cell.text = h
                cell.paragraphs[0].runs[0].font.bold = True

            total_general_ht = 0.0
            for r, item in enumerate(items):
                desig = str(item.get("designation", ""))
                try:
                    qte = float(item.get("qte", 0) or 0)
                except (TypeError, ValueError):
                    qte = 0.0
                try:
                    pu = float(item.get("pu", 0) or 0)
                except (TypeError, ValueError):
                    pu = 0.0
                row_total = qte * pu
                total_general_ht += row_total

                tbl_item.rows[r + 1].cells[0].text = desig
                tbl_item.rows[r + 1].cells[1].text = f"{qte:g}"
                tbl_item.rows[r + 1].cells[2].text = f"{pu:.2f} €"
                tbl_item.rows[r + 1].cells[3].text = f"{row_total:.2f} €"

            total_row = tbl_item.rows[len(items) + 1]
            total_row.cells[0].text = "TOTAL GENERAL HT"
            total_row.cells[0].paragraphs[0].runs[0].font.bold = True
            total_row.cells[3].text = f"{total_general_ht:.2f} €"
            total_row.cells[3].paragraphs[0].runs[0].font.bold = True

        # --- Galerie photo + texte libre ---
        content = payload.get("content") or []
        if content:
            doc.add_heading("Photos & remarques complémentaires", level=1)
            pending_pair = []

            def flush_pair():
                if not pending_pair:
                    return
                grid_table = doc.add_table(rows=1, cols=2)
                grid_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                grid_table.autofit = False

                for col_idx, item in enumerate(pending_pair):
                    cell = grid_table.rows[0].cells[col_idx]
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

                    base_width = 2.3 if item["is_portrait"] else 2.9
                    width_spec = Inches(base_width * item.get("scale", 1.0))
                    p.add_run().add_picture(item["path"], width=width_spec)

                    if item["caption"]:
                        p_cap = cell.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        r_cap = p_cap.add_run(f"Figure : {item['caption']}")
                        r_cap.font.italic = True
                        r_cap.font.size = Pt(8.5)

                doc.add_paragraph()
                pending_pair.clear()

            for entry in content:
                if entry.get("type") == "photo":
                    idx = entry.get("photo_index")
                    if idx is None or idx < 0 or idx >= len(photos):
                        continue
                    raw_bytes = photos[idx]
                    jpeg_bytes, width, height = _prepare_image_bytes_for_jpeg(raw_bytes)
                    tmp_path = os.path.join(
                        tempfile.gettempdir(), f"photo_{os.getpid()}_{random.randint(0, 999999)}.jpg"
                    )
                    with open(tmp_path, "wb") as f:
                        f.write(jpeg_bytes)
                    temp_files.append(tmp_path)

                    pending_pair.append({
                        "path": tmp_path,
                        "caption": (entry.get("caption") or "").strip(),
                        "is_portrait": height > width,
                        "scale": float(entry.get("scale") or 1.0),
                    })
                    if len(pending_pair) == 2:
                        flush_pair()
                elif entry.get("type") == "text":
                    flush_pair()
                    text = (entry.get("text") or "").strip()
                    if text:
                        p_text = doc.add_paragraph()
                        p_text.add_run(text)

            flush_pair()

        # --- Signatures ---
        sig_payload = payload.get("signatures") or {}
        sig_technicien = (sig_payload.get("technicien") or {}).get("nom") or payload.get("technicien", "")
        sig_client = (sig_payload.get("client") or {}).get("nom") or payload.get("client", "")
        sig_exterieur = (sig_payload.get("exterieur") or {}).get("nom") or ""

        doc.add_paragraph()
        doc.add_heading("Signatures", level=1)

        sig_table = doc.add_table(rows=2, cols=3)
        sig_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        sig_table.style = 'Table Grid'

        columns = [
            ("Technicien", sig_technicien, (sig_payload.get("technicien") or {}).get("image")),
            ("Client", sig_client, (sig_payload.get("client") or {}).get("image")),
            ("Intervenant extérieur", sig_exterieur, (sig_payload.get("exterieur") or {}).get("image")),
        ]
        for col_idx, (label, name, _img) in enumerate(columns):
            header_cell = sig_table.rows[0].cells[col_idx]
            p_header = header_cell.paragraphs[0]
            p_header.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_header = p_header.add_run(f"{label}\n{name or '—'}")
            r_header.font.bold = True
            r_header.font.size = Pt(10)

        sig_table.rows[1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        sig_table.rows[1].height = Inches(0.9)

        for col_idx, (label, name, img_data) in enumerate(columns):
            img_bytes = _decode_data_url_or_b64(img_data)
            if not img_bytes:
                continue
            sig_png_path = os.path.join(
                tempfile.gettempdir(),
                f"signature_{slugify(label)}_{os.getpid()}_{random.randint(0, 999999)}.png"
            )
            with open(sig_png_path, "wb") as f:
                f.write(img_bytes)
            temp_files.append(sig_png_path)

            body_cell = sig_table.rows[1].cells[col_idx]
            p_body = body_cell.paragraphs[0]
            p_body.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_body = p_body.add_run()
            run_body.add_picture(sig_png_path, width=Inches(1.7))

        out_buf = io.BytesIO()
        doc.save(out_buf)
        return out_buf.getvalue()

    finally:
        for f in temp_files:
            if f and os.path.exists(f):
                try:
                    os.remove(f)
                except Exception:
                    pass
